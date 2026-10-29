#!/usr/bin/env python3
"""Génère un compte rendu de chantier Word (.docx) à partir d'un JSON.

Usage : python3 generer_cr.py cr.json CR_Chantier_AAAA-MM-JJ.docx
Format du JSON : voir references/format-json.md
"""
import json
import os
import sys

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError:
    sys.exit("python-docx manquant : pip install python-docx")

GRIS = "F2F2F2"
ORANGE = "FFF4E5"
ROUGE = RGBColor(0xC0, 0x39, 0x2B)
VERT = RGBColor(0x2E, 0x7D, 0x32)


def hex_rgb(h):
    h = (h or "#1F3A5F").lstrip("#")
    if len(h) != 6:
        h = "1F3A5F"
    return RGBColor.from_string(h.upper())


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def bottom_border(paragraph, color):
    pPr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), "8")
    b.set(qn("w:space"), "1")
    b.set(qn("w:color"), color)
    bdr.append(b)
    pPr.append(bdr)


def set_widths(table, widths):
    table.autofit = False
    tblPr = table._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    grid = table._tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(widths[i].twips)))
    for row in table.rows:
        for i, c in enumerate(row.cells):
            c.width = widths[i]


def add_field(run, instr):
    for tag, text in (("begin", None), (None, instr), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


class CR:
    def __init__(self, data, base_dir):
        self.d = data
        self.base = base_dir
        ch = data.get("charte", {})
        self.ch = ch
        self.color_hex = (ch.get("couleur") or "#1F3A5F").lstrip("#").upper()
        self.color = hex_rgb(self.color_hex)
        self.font = ch.get("police") or "Arial"
        self.doc = Document()
        self._setup()

    def _setup(self):
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        for side in ("left_margin", "right_margin"):
            setattr(sec, side, Cm(2))
        sec.top_margin, sec.bottom_margin = Cm(1.5), Cm(1.8)
        st = self.doc.styles["Normal"]
        st.font.name = self.font
        st.font.size = Pt(10)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), self.font)
        st.paragraph_format.space_after = Pt(3)

    def path(self, p):
        if not p:
            return None
        cand = p if os.path.isabs(p) else os.path.join(self.base, p)
        if os.path.exists(cand):
            return cand
        skill_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        cand2 = os.path.join(skill_root, p)
        return cand2 if os.path.exists(cand2) else None

    # ---------- blocs ----------
    def entete(self):
        d, ch = self.d, self.ch
        t = self.doc.add_table(rows=1, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_widths(t, [Cm(9), Cm(8)])
        left, right = t.rows[0].cells
        logo = self.path(ch.get("logo"))
        if logo:
            left.paragraphs[0].add_run().add_picture(logo, height=Cm(1.6))
        else:
            r = left.paragraphs[0].add_run(ch.get("entreprise") or "[Entreprise]")
            r.bold, r.font.size, r.font.color.rgb = True, Pt(13), self.color
        if ch.get("adresse"):
            p = left.add_paragraph(ch["adresse"])
            p.runs[0].font.size = Pt(8)
        p = right.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run("COMPTE RENDU DE CHANTIER")
        r.bold, r.font.size, r.font.color.rgb = True, Pt(14), self.color
        num = d.get("numero")
        p2 = right.add_paragraph(f"CR n° {num if num not in (None, '') else '[n°]'}")
        p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p2.runs[0].bold = True

        self.doc.add_paragraph()
        rows = [
            ("Chantier", " — ".join(x for x in [d.get("chantier"), d.get("adresse_chantier")] if x)),
            ("Type", d.get("type") or "Visite de chantier"),
            ("Date", " à ".join(x for x in [d.get("date"), d.get("heure")] if x)),
            ("Rédacteur", d.get("redacteur")),
            ("Météo", d.get("meteo")),
            ("Présents", ", ".join(d.get("presents") or [])),
            ("Absents excusés", ", ".join(d.get("absents_excuses") or [])),
            ("Diffusion", ", ".join(d.get("diffusion") or [])),
        ]
        rows = [(k, v) for k, v in rows if v or k in ("Chantier", "Date", "Présents")]
        tb = self.doc.add_table(rows=len(rows), cols=2)
        tb.style = "Table Grid"
        for i, (k, v) in enumerate(rows):
            a, b = tb.rows[i].cells
            a.width, b.width = Cm(4), Cm(13)
            shade(a, GRIS)
            a.paragraphs[0].add_run(k).bold = True
            b.paragraphs[0].add_run(v or "Néant")
        set_widths(tb, [Cm(4), Cm(13)])

    def titre(self, n, txt):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(f"{n}. {txt.upper()}")
        r.bold, r.font.size, r.font.color.rgb = True, Pt(11), self.color
        bottom_border(p, self.color_hex)

    def neant(self):
        self.doc.add_paragraph("Néant").runs[0].italic = True

    def liste(self, items):
        if not items:
            return self.neant()
        for it in items:
            self.doc.add_paragraph(str(it), style="List Bullet")

    def avancement(self, items):
        if not items:
            return self.neant()
        for it in items:
            if isinstance(it, str):
                self.doc.add_paragraph(it, style="List Bullet")
                continue
            p = self.doc.add_paragraph(style="List Bullet")
            if it.get("lot"):
                p.add_run(f"{it['lot']} : ").bold = True
            p.add_run(it.get("texte", ""))

    def bloquants(self, items):
        if not items:
            return self.neant()
        for it in items:
            p = self.doc.add_paragraph()
            if isinstance(it, str):
                p.add_run(it).bold = True
                continue
            r = p.add_run(f"{it.get('titre', 'Point bloquant')} — ")
            r.bold, r.font.color.rgb = True, ROUGE
            p.add_run(it.get("texte", ""))

    def actions(self, items):
        if not items:
            return self.neant()
        hdr = ["N°", "Action", "Responsable", "Échéance", "Statut"]
        widths = [Cm(1), Cm(7.5), Cm(3.2), Cm(2.5), Cm(2.8)]
        t = self.doc.add_table(rows=1, cols=5)
        t.style = "Table Grid"
        for i, h in enumerate(hdr):
            c = t.rows[0].cells[i]
            c.width = widths[i]
            shade(c, self.color_hex)
            r = c.paragraphs[0].add_run(h)
            r.bold, r.font.color.rgb = True, RGBColor(0xFF, 0xFF, 0xFF)
        for n, a in enumerate(items, 1):
            cells = t.add_row().cells
            vals = [str(a.get("n", n)), a.get("action", ""), a.get("responsable") or "⚠ à confirmer",
                    a.get("echeance") or "⚠ à confirmer", a.get("statut") or "À faire"]
            for i, v in enumerate(vals):
                cells[i].width = widths[i]
                r = cells[i].paragraphs[0].add_run(v)
                if i == 4:
                    low = v.lower()
                    if "retard" in low:
                        r.font.color.rgb, r.bold = ROUGE, True
                    elif "sold" in low:
                        r.font.color.rgb = VERT
                if "⚠" in v:
                    r.font.color.rgb = ROUGE
        set_widths(t, widths)

    def a_confirmer(self, items):
        if not items:
            return self.neant()
        t = self.doc.add_table(rows=1, cols=1)
        t.style = "Table Grid"
        c = t.rows[0].cells[0]
        shade(c, ORANGE)
        first = True
        for n, it in enumerate(items, 1):
            p = c.paragraphs[0] if first else c.add_paragraph()
            first = False
            if isinstance(it, str):
                p.add_run(f"⚠ {n}. {it}")
                continue
            p.add_run(f"⚠ {n}. {it.get('point', '')}").bold = True
            if it.get("extrait"):
                r = p.add_run(f" — dicté : {it['extrait']}")
                r.italic = True

    def photos(self, items):
        ok = [(self.path(it.get("fichier")), it.get("legende", "")) for it in items or []]
        for (pth, _), it in zip(ok, items or []):
            if not pth:
                print(f"Photo introuvable, ignorée : {it.get('fichier')}", file=sys.stderr)
        ok = [x for x in ok if x[0]]
        if not ok:
            return False
        self.doc.add_page_break()
        p = self.doc.add_paragraph()
        r = p.add_run("ANNEXE — PHOTOS")
        r.bold, r.font.size, r.font.color.rgb = True, Pt(11), self.color
        bottom_border(p, self.color_hex)
        for n, (pth, leg) in enumerate(ok, 1):
            self.doc.add_paragraph().add_run().add_picture(pth, width=Cm(12))
            cap = self.doc.add_paragraph(f"Photo {n} — {leg}")
            cap.runs[0].italic, cap.runs[0].font.size = True, Pt(9)
        return True

    def pied(self):
        d, ch = self.d, self.ch
        sec = self.doc.sections[0]
        p = sec.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        num = d.get("numero")
        r = p.add_run(f"{ch.get('entreprise') or '[Entreprise]'} · CR n° {num if num not in (None, '') else '[n°]'} — "
                      f"{d.get('chantier', '')} · Page ")
        r.font.size = Pt(8)
        r2 = p.add_run()
        r2.font.size = Pt(8)
        add_field(r2, "PAGE")
        r3 = p.add_run(" / ")
        r3.font.size = Pt(8)
        r4 = p.add_run()
        r4.font.size = Pt(8)
        add_field(r4, "NUMPAGES")

    def build(self, out):
        d = self.d
        self.entete()
        sections = [
            ("Avancement des travaux", lambda: self.avancement(d.get("avancement"))),
            ("Points bloquants et alertes", lambda: self.bloquants(d.get("bloquants"))),
            ("Planning", lambda: self.liste(d.get("planning"))),
            ("Sous-traitants", lambda: self.liste(d.get("sous_traitants"))),
            ("Sécurité, propreté, incidents", lambda: self.liste(d.get("securite"))),
            ("Approvisionnements et livraisons", lambda: self.liste(d.get("livraisons"))),
            ("Décisions prises", lambda: self.liste(d.get("decisions"))),
            ("Actions à mener", lambda: self.actions(d.get("actions"))),
            ("Prochaine visite / réunion", lambda: self.liste([d["prochaine_visite"]] if d.get("prochaine_visite") else [])),
            ("Points à confirmer", lambda: self.a_confirmer(d.get("a_confirmer"))),
        ]
        for i, (t, fn) in enumerate(sections, 1):
            self.titre(i, t)
            fn()
        if self.ch.get("delai_observations"):
            self.doc.add_paragraph()
            p = self.doc.add_paragraph(
                f"Sauf observations écrites dans un délai de {self.ch['delai_observations']} à compter de sa "
                "diffusion, le présent compte rendu est réputé accepté par l'ensemble des destinataires.")
            p.runs[0].italic, p.runs[0].font.size = True, Pt(8.5)
        self.photos(d.get("photos"))
        self.pied()
        self.doc.save(out)
        print(f"CR généré : {out}")


def main():
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    src, out = sys.argv[1], sys.argv[2]
    with open(src, encoding="utf-8") as f:
        data = json.load(f)
    if not data.get("chantier") or not data.get("date"):
        sys.exit("Champs obligatoires manquants : 'chantier' et 'date'.")
    CR(data, os.path.dirname(os.path.abspath(src))).build(out)


if __name__ == "__main__":
    main()
